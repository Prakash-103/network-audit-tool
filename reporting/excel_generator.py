"""
Excel Audit Evidence Workbook Generator & Importer.
Creates & Parses a clean 3-sheet .xlsx workbook matching the 3-Stage UI workflow:
Sheet 1: Device Inventory (Stage 1)
Sheet 2: Link Connections (Stage 2)
Sheet 3: Audit Status (Stage 3)
"""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models import (
    NetworkProjectState, DeviceModel, TopologyLink, PortChannelGroup, DeviceAuditTask
)
from core.constants import (
    DeviceFunction, DeviceRole, LifecycleStatus, TaskVerificationStatus, AuditStatus, SeverityLevel
)
from engine.protocol_catalog import ProtocolCatalogEngine

class ExcelWorkbookGenerator:
    @staticmethod
    def generate(state: NetworkProjectState, output_path: str) -> str:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        # High-contrast styling
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        section_font = Font(name="Segoe UI", size=12, bold=True, color="1E3A8A")
        
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        section_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        status_fills = {
            "Completed": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "PASS": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "Verified": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            
            "Failed": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
            "FAIL": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
            
            "Warning": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
            "WARNING": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
            
            "In-Complete": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
            "N/A": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
        }

        def write_headers(ws, row_idx, headers):
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_num, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        def autofit(ws):
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(min(max_len + 4, 60), 12)

        # =========================================================================
        # SHEET 1: DEVICE INVENTORY (Stage 1)
        # =========================================================================
        ws1 = wb.create_sheet(title="Device Inventory")
        ws1.views.sheetView[0].showGridLines = True
        
        headers1 = [
            "Device ID", "Hostname", "Company Name", "Vendor", "Hardware Model", "OS / Version",
            "Function", "Role", "Ports", "Port Speed", "Management IP",
            "Site Location", "Serial Number", "Lifecycle Status"
        ]
        write_headers(ws1, 1, headers1)

        for dev in state.devices:
            ws1.append([
                dev.device_id,
                dev.hostname,
                dev.company_name,
                dev.vendor,
                dev.device_model,
                dev.os_version,
                dev.device_function.value if hasattr(dev.device_function, 'value') else dev.device_function,
                dev.device_role.value if hasattr(dev.device_role, 'value') else dev.device_role,
                dev.num_ports,
                dev.port_speed,
                dev.management_ip,
                dev.site_location,
                dev.serial_number,
                dev.lifecycle_status.value if hasattr(dev.lifecycle_status, 'value') else dev.lifecycle_status
            ])
        autofit(ws1)

        # =========================================================================
        # SHEET 2: LINK CONNECTIONS (Stage 2)
        # =========================================================================
        ws2 = wb.create_sheet(title="Link Connections")
        ws2.views.sheetView[0].showGridLines = True

        # Section 1: Physical Device Links
        headers2_1 = ["Source Device", "Source Port", "Target Device", "Target Port", "Link Speed", "Port-Channel", "Status"]
        write_headers(ws2, 1, headers2_1)

        for link in state.topology_links:
            ws2.append([
                link.source_device,
                link.source_port,
                link.target_device,
                link.target_port,
                link.link_speed,
                link.port_channel_id or "Single Link",
                link.status
            ])

        # Blank spacing row
        curr_row = ws2.max_row + 3
        
        # Section 2: Port-Channel Groups Header
        cell_sec2 = ws2.cell(row=curr_row, column=1, value="Port-Channel Aggregations")
        cell_sec2.font = section_font
        curr_row += 1

        headers2_2 = ["Port-Channel ID", "Source Device", "Target Device", "Source Member Ports", "Target Member Ports", "Protocol", "Status"]
        write_headers(ws2, curr_row, headers2_2)

        for po in state.port_channels:
            ws2.append([
                po.id,
                po.source_device,
                po.target_device,
                ", ".join(po.source_member_ports),
                ", ".join(po.target_member_ports),
                po.protocol,
                po.status
            ])

        # Section 3: Topology Validation Alerts
        if state.validation_issues:
            curr_row = ws2.max_row + 3
            cell_sec3 = ws2.cell(row=curr_row, column=1, value="Real-Time Topology Validation Issues")
            cell_sec3.font = section_font
            curr_row += 1

            headers2_3 = ["Rule", "Severity", "Source Device", "Issue Message", "Remediation"]
            write_headers(ws2, curr_row, headers2_3)

            for issue in state.validation_issues:
                ws2.append([
                    issue.rule,
                    issue.severity.value if hasattr(issue.severity, 'value') else issue.severity,
                    issue.source_device,
                    issue.message,
                    issue.remediation
                ])

        autofit(ws2)

        # =========================================================================
        # SHEET 3: AUDIT STATUS (Stage 3)
        # =========================================================================
        ws3 = wb.create_sheet(title="Audit Status")
        ws3.views.sheetView[0].showGridLines = True

        headers3 = [
            "Task ID", "Device Hostname", "Device Classification", "Role",
            "Category", "Protocol", "Title / Test", "Condition",
            "CLI Verification Command", "Verification Status", "Evidence Notes", "Actual Output / Findings"
        ]
        write_headers(ws3, 1, headers3)

        dev_map = {d.hostname: d for d in state.devices}

        for task in state.device_audit_tasks:
            dev = dev_map.get(task.device_hostname)
            classification = ProtocolCatalogEngine.detect_device_classification(dev) if dev else "Network Device"
            
            row_vals = [
                task.id,
                task.device_hostname,
                classification,
                task.device_role,
                task.category,
                task.protocol,
                task.title,
                task.condition,
                task.verification_command,
                task.status.value if hasattr(task.status, 'value') else task.status,
                task.evidence_notes,
                task.actual_output
            ]
            ws3.append(row_vals)

            # Style status cell
            status_str = task.status.value if hasattr(task.status, 'value') else str(task.status)
            if status_str in status_fills:
                ws3.cell(row=ws3.max_row, column=10).fill = status_fills[status_str]

        autofit(ws3)

        wb.save(output_path)
        return output_path


class ExcelWorkbookImporter:
    """
    Parses an uploaded 3-sheet Excel workbook (Device Inventory, Link Connections, Audit Status)
    and reconstructs the NetworkProjectState.
    """
    @staticmethod
    def parse_workbook(file_bytes: bytes) -> NetworkProjectState:
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
        sheet_names = wb.sheetnames

        devices = []
        links = []
        port_channels = []
        audit_tasks = []

        # Map enum helper
        def match_enum(enum_class, value: str, default_val):
            if not value:
                return default_val
            val_str = str(value).strip().lower()
            for member in enum_class:
                if member.value.lower() == val_str or member.name.lower() == val_str:
                    return member
            return default_val

        # ----------------------------------------------------
        # 1. PARSE SHEET 1: DEVICE INVENTORY
        # ----------------------------------------------------
        ws1 = None
        for name in sheet_names:
            if "inventory" in name.lower() or "device" in name.lower():
                ws1 = wb[name]
                break

        if ws1:
            for r in range(2, ws1.max_row + 1):
                row = [ws1.cell(row=r, column=c).value for c in range(1, 15)]
                if not row[1] and not row[0]:  # No hostname or ID
                    continue
                dev_id = str(row[0] or f"DEV-{r-1}").strip()
                hostname = str(row[1] or f"NODE-{r-1}").strip()
                comp_name = str(row[2] or "Enterprise Org").strip()
                vendor = str(row[3] or "Cisco").strip()
                model = str(row[4] or "Catalyst 9300").strip()
                os_ver = str(row[5] or "IOS-XE 17.9").strip()
                func_val = match_enum(DeviceFunction, row[6], DeviceFunction.L2_L3)
                role_val = match_enum(DeviceRole, row[7], DeviceRole.ACCESS)
                
                try:
                    num_ports = int(row[8]) if row[8] is not None else 24
                except (ValueError, TypeError):
                    num_ports = 24

                speed = str(row[9] or "1G").strip()
                mgmt_ip = str(row[10] or "10.254.1.1").strip()
                site = str(row[11] or "HQ DC").strip()
                serial = str(row[12] or f"SN-{r-1}").strip()
                lifecycle_val = match_enum(LifecycleStatus, row[13], LifecycleStatus.ACTIVE)

                dev = DeviceModel(
                    device_id=dev_id,
                    hostname=hostname,
                    company_name=comp_name,
                    vendor=vendor,
                    device_model=model,
                    os_version=os_ver,
                    device_function=func_val,
                    device_role=role_val,
                    num_ports=num_ports,
                    port_speed=speed,
                    management_ip=mgmt_ip,
                    site_location=site,
                    serial_number=serial,
                    lifecycle_status=lifecycle_val,
                    ports=[]
                )
                devices.append(dev)

        # ----------------------------------------------------
        # 2. PARSE SHEET 2: LINK CONNECTIONS
        # ----------------------------------------------------
        ws2 = None
        for name in sheet_names:
            if "link" in name.lower() or "topology" in name.lower() or "connection" in name.lower():
                ws2 = wb[name]
                break

        if ws2:
            section = "LINKS"
            for r in range(2, ws2.max_row + 1):
                col1 = ws2.cell(row=r, column=1).value
                col1_str = str(col1 or "").strip()

                if "port-channel aggregations" in col1_str.lower():
                    section = "PORT_CHANNELS"
                    continue
                if "topology validation" in col1_str.lower():
                    section = "VALIDATION"
                    continue
                if col1_str.lower() in ["source device", "port-channel id", "rule"]:
                    continue  # Skip section headers

                if section == "LINKS":
                    src_dev = col1_str
                    src_port = str(ws2.cell(row=r, column=2).value or "").strip()
                    tgt_dev = str(ws2.cell(row=r, column=3).value or "").strip()
                    tgt_port = str(ws2.cell(row=r, column=4).value or "").strip()
                    speed = str(ws2.cell(row=r, column=5).value or "10G").strip()
                    po_id = str(ws2.cell(row=r, column=6).value or "").strip()
                    status = str(ws2.cell(row=r, column=7).value or "Valid").strip()

                    if src_dev and src_port and tgt_dev and tgt_port:
                        if po_id.lower() in ["single link", "none", "n/a", ""]:
                            po_id = None
                        link_id = f"{src_dev}_{src_port}___{tgt_dev}_{tgt_port}"
                        link = TopologyLink(
                            id=link_id,
                            source_device=src_dev,
                            source_port=src_port,
                            target_device=tgt_dev,
                            target_port=tgt_port,
                            link_speed=speed,
                            port_channel_id=po_id,
                            status=status
                        )
                        links.append(link)

                elif section == "PORT_CHANNELS":
                    po_id = col1_str
                    src_dev = str(ws2.cell(row=r, column=2).value or "").strip()
                    tgt_dev = str(ws2.cell(row=r, column=3).value or "").strip()
                    src_ports_raw = str(ws2.cell(row=r, column=4).value or "").strip()
                    tgt_ports_raw = str(ws2.cell(row=r, column=5).value or "").strip()
                    protocol = str(ws2.cell(row=r, column=6).value or "LACP").strip()
                    status = str(ws2.cell(row=r, column=7).value or "Operational").strip()

                    if po_id and src_dev and tgt_dev:
                        src_ports = [p.strip() for p in src_ports_raw.split(",") if p.strip()]
                        tgt_ports = [p.strip() for p in tgt_ports_raw.split(",") if p.strip()]
                        po_group = PortChannelGroup(
                            id=po_id,
                            source_device=src_dev,
                            target_device=tgt_dev,
                            source_member_ports=src_ports,
                            target_member_ports=tgt_ports,
                            protocol=protocol,
                            status=status
                        )
                        port_channels.append(po_group)

        # ----------------------------------------------------
        # 3. PARSE SHEET 3: AUDIT STATUS
        # ----------------------------------------------------
        ws3 = None
        for name in sheet_names:
            if "audit" in name.lower() or "task" in name.lower() or "status" in name.lower():
                ws3 = wb[name]
                break

        if ws3:
            for r in range(2, ws3.max_row + 1):
                task_id = str(ws3.cell(row=r, column=1).value or "").strip()
                dev_host = str(ws3.cell(row=r, column=2).value or "").strip()
                dev_role = str(ws3.cell(row=r, column=4).value or "Access").strip()
                cat = str(ws3.cell(row=r, column=5).value or "General").strip()
                proto = str(ws3.cell(row=r, column=6).value or "Health").strip()
                title = str(ws3.cell(row=r, column=7).value or "Audit Test").strip()
                cond = str(ws3.cell(row=r, column=8).value or "Always Applicable").strip()
                cli_cmd = str(ws3.cell(row=r, column=9).value or "").strip()
                status_raw = str(ws3.cell(row=r, column=10).value or "In-Complete").strip()
                notes = str(ws3.cell(row=r, column=11).value or "").strip()
                output = str(ws3.cell(row=r, column=12).value or "").strip()

                if task_id and dev_host:
                    status_enum = match_enum(TaskVerificationStatus, status_raw, TaskVerificationStatus.INCOMPLETE)
                    task = DeviceAuditTask(
                        id=task_id,
                        device_hostname=dev_host,
                        device_role=dev_role,
                        protocol=proto,
                        category=cat,
                        title=title,
                        description="",
                        condition=cond,
                        verification_command=cli_cmd,
                        status=status_enum,
                        actual_output=output,
                        evidence_notes=notes
                    )
                    audit_tasks.append(task)

        comp_name = devices[0].company_name if devices else "Imported Infrastructure Audit"
        site_name = devices[0].site_location if devices else "HQ Data Center"

        new_state = NetworkProjectState(
            project_name=f"{comp_name} - {site_name} Audit",
            company_name=comp_name,
            site_location=site_name,
            devices=devices,
            topology_links=links,
            port_channels=port_channels,
            device_audit_tasks=audit_tasks
        )

        return new_state


